from pathlib import Path
import json
import math
import shutil
import zipfile

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    Image as ReportLabImage,
)

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font


# ============================================================
# CONFIGURATION
# ============================================================

BASE_DIR = Path(__file__).parent
OUTPUT_DIR = BASE_DIR / "rag_test_dataset"

OUTPUT_DIR.mkdir(exist_ok=True)


# ============================================================
# HELPER FUNCTIONS
# ============================================================

def get_font(size=24, bold=False):
    """Load a Windows-friendly font."""
    candidates = []

    if bold:
        candidates = [
            r"C:\Windows\Fonts\arialbd.ttf",
            r"C:\Windows\Fonts\calibrib.ttf",
        ]
    else:
        candidates = [
            r"C:\Windows\Fonts\arial.ttf",
            r"C:\Windows\Fonts\calibri.ttf",
        ]

    for font_path in candidates:
        if Path(font_path).exists():
            return ImageFont.truetype(font_path, size)

    return ImageFont.load_default()


def draw_centered_text(draw, box, text, font, fill="black"):
    x1, y1, x2, y2 = box

    bbox = draw.textbbox((0, 0), text, font=font)

    width = bbox[2] - bbox[0]
    height = bbox[3] - bbox[1]

    x = x1 + ((x2 - x1) - width) / 2
    y = y1 + ((y2 - y1) - height) / 2

    draw.text((x, y), text, font=font, fill=fill)


def draw_arrow(draw, start, end):
    x1, y1 = start
    x2, y2 = end

    draw.line(
        (x1, y1, x2, y2),
        fill="black",
        width=5
    )

    angle = math.atan2(y2 - y1, x2 - x1)

    arrow_length = 22
    arrow_angle = 0.5

    p1 = (
        x2 - arrow_length * math.cos(angle - arrow_angle),
        y2 - arrow_length * math.sin(angle - arrow_angle),
    )

    p2 = (
        x2 - arrow_length * math.cos(angle + arrow_angle),
        y2 - arrow_length * math.sin(angle + arrow_angle),
    )

    draw.polygon(
        [(x2, y2), p1, p2],
        fill="black"
    )


# ============================================================
# 1. ARCHITECTURE PNG
# ============================================================

def create_architecture_png():

    path = OUTPUT_DIR / "architecture.png"

    width = 1400
    height = 800

    image = Image.new(
        "RGB",
        (width, height),
        "white"
    )

    draw = ImageDraw.Draw(image)

    title_font = get_font(40, bold=True)
    box_font = get_font(28, bold=True)
    normal_font = get_font(24)

    draw.text(
        (390, 40),
        "NovaTech Application Architecture",
        font=title_font,
        fill="black",
    )

    boxes = {
        "Customer": (100, 300, 350, 420),
        "API Gateway": (550, 280, 850, 440),
        "Payment Service": (1050, 170, 1320, 290),
        "Order Service": (1050, 500, 1320, 620),
    }

    for label, box in boxes.items():

        draw.rounded_rectangle(
            box,
            radius=25,
            outline="black",
            width=4,
        )

        draw_centered_text(
            draw,
            box,
            label,
            box_font,
        )

    # Customer -> Gateway
    draw_arrow(
        draw,
        (350, 360),
        (550, 360),
    )

    # Gateway -> Payment
    draw_arrow(
        draw,
        (850, 330),
        (1050, 230),
    )

    # Gateway -> Order
    draw_arrow(
        draw,
        (850, 390),
        (1050, 560),
    )

    draw.text(
        (370, 700),
        "Customer → API Gateway → Payment Service / Order Service",
        font=normal_font,
        fill="black",
    )

    image.save(path)

    print(f"Created: {path}")


# ============================================================
# 2. REMOTE WORK IMAGE
# ============================================================

def create_remote_work_image():

    path = OUTPUT_DIR / "remote_work_diagram.png"

    width = 1000
    height = 500

    image = Image.new(
        "RGB",
        (width, height),
        "white"
    )

    draw = ImageDraw.Draw(image)

    title_font = get_font(36, bold=True)
    box_font = get_font(28, bold=True)
    normal_font = get_font(24)

    draw.text(
        (320, 30),
        "Remote Work Model",
        font=title_font,
        fill="black",
    )

    office_box = (100, 180, 350, 310)
    remote_box = (650, 180, 900, 310)

    for label, box in [
        ("Office", office_box),
        ("Remote", remote_box),
    ]:

        draw.rounded_rectangle(
            box,
            radius=20,
            outline="black",
            width=4,
        )

        draw_centered_text(
            draw,
            box,
            label,
            box_font,
        )

    draw_arrow(
        draw,
        (350, 245),
        (650, 245),
    )

    draw.text(
        (285, 370),
        "Up to 2 remote days per week",
        font=normal_font,
        fill="black",
    )

    image.save(path)

    print(f"Created: {path}")


# ============================================================
# 3. EMPLOYEE HANDBOOK DOCX
# ============================================================

def create_employee_handbook():

    path = OUTPUT_DIR / "employee_handbook.docx"

    doc = Document()

    doc.add_heading(
        "NovaTech Employee Handbook",
        level=0
    )

    doc.add_paragraph(
        "This handbook summarizes key workplace policies "
        "for full-time employees of NovaTech Solutions."
    )

    # Leave
    doc.add_heading(
        "Leave Policy",
        level=1
    )

    doc.add_paragraph(
        "All full-time employees are entitled to "
        "24 days of paid annual leave per calendar year."
    )

    doc.add_paragraph(
        "Leave requests must be submitted at least "
        "7 days before the intended start date."
    )

    # Working hours
    doc.add_heading(
        "Working Hours",
        level=1
    )

    doc.add_paragraph(
        "Standard working hours are 9:00 AM to 6:00 PM, "
        "Monday through Friday."
    )

    # Remote work
    doc.add_heading(
        "Remote Work",
        level=1
    )

    doc.add_paragraph(
        "Employees may work remotely up to 2 days per week, "
        "subject to manager approval."
    )

    # Embedded image
    doc.add_picture(
        str(OUTPUT_DIR / "remote_work_diagram.png"),
        width=Inches(5.8),
    )

    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(
        "Figure 1: NovaTech remote work model."
    )

    # Table
    doc.add_heading(
        "Benefits Table",
        level=1
    )

    table = doc.add_table(
        rows=1,
        cols=3,
    )

    headers = [
        "Benefit",
        "Eligibility",
        "Annual Allowance",
    ]

    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header

    rows = [
        (
            "Paid Annual Leave",
            "Full-time employees",
            "24 days",
        ),
        (
            "Remote Work",
            "Manager approval",
            "Up to 2 days/week",
        ),
        (
            "Training Leave",
            "After probation",
            "5 days",
        ),
    ]

    for row in rows:

        cells = table.add_row().cells

        for i, value in enumerate(row):
            cells[i].text = value

    # Contact
    doc.add_heading(
        "Important Contact",
        level=1
    )

    doc.add_paragraph(
        "HR Operations: hr@novatech.example"
    )

    doc.save(path)

    print(f"Created: {path}")


# ============================================================
# 4. PAYMENT POLICY PDF
# ============================================================

def create_payment_policy():

    path = OUTPUT_DIR / "payment_policy.pdf"

    styles = getSampleStyleSheet()

    document = SimpleDocTemplate(
        str(path),
        pagesize=letter,
    )

    story = []

    story.append(
        Paragraph(
            "NovaTech Payment Policy",
            styles["Title"],
        )
    )

    story.append(Spacer(1, 15))

    story.append(
        Paragraph(
            "The Payment API and internal payment operations "
            "follow the limits described in this document.",
            styles["BodyText"],
        )
    )

    story.append(Spacer(1, 15))

    story.append(
        Paragraph(
            "Payment Limits",
            styles["Heading2"],
        )
    )

    story.append(
        Paragraph(
            "The minimum standard payment amount is "
            "<b>$1.00</b>. The maximum standard payment "
            "amount is <b>$10,000.00 per transaction</b>.",
            styles["BodyText"],
        )
    )

    story.append(Spacer(1, 15))

    data = [
        [
            "Payment Type",
            "Minimum",
            "Maximum",
            "Currency",
        ],
        [
            "Standard transaction",
            "$1.00",
            "$10,000.00",
            "USD",
        ],
        [
            "Refund",
            "$1.00",
            "$5,000.00",
            "USD",
        ],
        [
            "Manual adjustment",
            "$1.00",
            "$2,500.00",
            "USD",
        ],
    ]

    table = Table(
        data,
        colWidths=[
            150,
            100,
            110,
            80,
        ],
    )

    table.setStyle(
        TableStyle([
            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.8,
                colors.black,
            ),
            (
                "BACKGROUND",
                (0, 0),
                (-1, 0),
                colors.lightgrey,
            ),
            (
                "FONTNAME",
                (0, 0),
                (-1, 0),
                "Helvetica-Bold",
            ),
        ])
    )

    story.append(table)

    story.append(Spacer(1, 20))

    story.append(
        Paragraph(
            "Authentication",
            styles["Heading2"],
        )
    )

    story.append(
        Paragraph(
            "The Payment API uses <b>Bearer Token authentication</b>. "
            "Clients must send the token in the HTTP Authorization header.",
            styles["BodyText"],
        )
    )

    story.append(Spacer(1, 20))

    story.append(
        Paragraph(
            "Architecture Reference",
            styles["Heading2"],
        )
    )

    story.append(
        Paragraph(
            "The following diagram shows how customer requests "
            "reach the API Gateway and are routed to the Payment "
            "Service or Order Service.",
            styles["BodyText"],
        )
    )

    story.append(Spacer(1, 15))

    story.append(
        ReportLabImage(
            str(OUTPUT_DIR / "architecture.png"),
            width=500,
            height=286,
        )
    )

    document.build(story)

    print(f"Created: {path}")


# ============================================================
# 5. CSV
# ============================================================

def create_product_catalog():

    path = OUTPUT_DIR / "product_catalog.csv"

    content = """product_id,product_name,category,price,currency,stock
P001,Wireless Keyboard,Accessories,49.99,USD,120
P002,Wireless Mouse,Accessories,29.99,USD,250
P003,USB-C Hub,Accessories,39.99,USD,85
P004,27-inch Monitor,Monitors,249.99,USD,45
P005,Mechanical Keyboard,Accessories,89.99,USD,60
P006,Laptop Stand,Office,34.99,USD,150
"""

    path.write_text(
        content,
        encoding="utf-8",
    )

    print(f"Created: {path}")


# ============================================================
# 6. JSON
# ============================================================

def create_customers():

    path = OUTPUT_DIR / "customers.json"

    data = {
        "customers": [
            {
                "customer_id": "C1001",
                "name": "Alice Johnson",
                "membership": "Gold",
                "country": "India",
                "orders": 18,
                "total_spent": 2450.75,
            },
            {
                "customer_id": "C1002",
                "name": "Bob Smith",
                "membership": "Silver",
                "country": "United States",
                "orders": 9,
                "total_spent": 980.50,
            },
            {
                "customer_id": "C1003",
                "name": "Carol Williams",
                "membership": "Platinum",
                "country": "United Kingdom",
                "orders": 32,
                "total_spent": 5620.25,
            },
            {
                "customer_id": "C1004",
                "name": "David Brown",
                "membership": "Bronze",
                "country": "Australia",
                "orders": 4,
                "total_spent": 320.00,
            },
        ]
    }

    path.write_text(
        json.dumps(
            data,
            indent=2,
        ),
        encoding="utf-8",
    )

    print(f"Created: {path}")


# ============================================================
# 7. HTML
# ============================================================

def create_company_html():

    path = OUTPUT_DIR / "company_info.html"

    html = """<!DOCTYPE html>
<html>
<head>
    <meta charset="UTF-8">
    <title>NovaTech Solutions</title>
</head>

<body>

<h1>NovaTech Solutions</h1>

<h2>About Us</h2>

<p>
NovaTech Solutions is a software company specializing in
cloud infrastructure and data analytics.
</p>

<h2>Company Details</h2>

<ul>
    <li>Founded: 2018</li>
    <li>Headquarters: Bengaluru, India</li>
    <li>Employees: 450</li>
    <li>Industry: Software and Technology</li>
</ul>

<h2>Products</h2>

<p>
Our primary products include NovaCloud, NovaAnalytics,
and NovaSecure.
</p>

<h2>Architecture</h2>

<p>
The architecture diagram shows how customers connect to
the API Gateway and how requests are routed to the
Payment Service and Order Service.
</p>

<img
    src="architecture.png"
    alt="NovaTech application architecture diagram"
    width="900"
/>

</body>
</html>
"""

    path.write_text(
        html,
        encoding="utf-8",
    )

    print(f"Created: {path}")


# ============================================================
# 8. EXCEL
# ============================================================

def create_financial_report():

    path = OUTPUT_DIR / "financial_report.xlsx"

    workbook = Workbook()

    # --------------------------------------------------------
    # Revenue sheet
    # --------------------------------------------------------

    sheet = workbook.active

    sheet.title = "Revenue"

    sheet.append([
        "Year",
        "Revenue_USD_Millions",
    ])

    revenue = [
        (2022, 80),
        (2023, 100),
        (2024, 130),
        (2025, 150),
    ]

    for row in revenue:
        sheet.append(row)

    for cell in sheet[1]:
        cell.font = Font(
            bold=True
        )

    chart = BarChart()

    chart.title = "Annual Revenue"

    chart.y_axis.title = "USD Millions"

    chart.x_axis.title = "Year"

    data = Reference(
        sheet,
        min_col=2,
        min_row=1,
        max_row=5,
    )

    categories = Reference(
        sheet,
        min_col=1,
        min_row=2,
        max_row=5,
    )

    chart.add_data(
        data,
        titles_from_data=True,
    )

    chart.set_categories(
        categories
    )

    chart.height = 8

    chart.width = 14

    sheet.add_chart(
        chart,
        "D2",
    )

    # --------------------------------------------------------
    # Department Costs sheet
    # --------------------------------------------------------

    costs = workbook.create_sheet(
        "Department Costs"
    )

    costs.append([
        "Department",
        "2025 Cost_USD_M",
    ])

    department_data = [
        ("Engineering", 4.2),
        ("Finance", 2.1),
        ("Operations", 3.5),
        ("Sales", 2.8),
    ]

    for row in department_data:
        costs.append(row)

    for cell in costs[1]:
        cell.font = Font(
            bold=True
        )

    workbook.save(path)

    print(f"Created: {path}")


# ============================================================
# 9. RAG QUESTIONS
# ============================================================

def create_questions():

    path = OUTPUT_DIR / "rag_test_questions.json"

    questions = [

        {
            "id": "Q01",
            "question": "How many paid annual leave days do full-time employees receive?",
            "expected_answer": "Full-time employees receive 24 days of paid annual leave per calendar year.",
            "source": "employee_handbook.docx",
            "type": "direct"
        },

        {
            "id": "Q02",
            "question": "How many days per week can an employee work remotely?",
            "expected_answer": "Employees may work remotely up to 2 days per week, subject to manager approval.",
            "source": "employee_handbook.docx",
            "type": "table/text"
        },

        {
            "id": "Q03",
            "question": "What is the annual allowance for training leave?",
            "expected_answer": "The annual training leave allowance is 5 days after probation.",
            "source": "employee_handbook.docx",
            "type": "table"
        },

        {
            "id": "Q04",
            "question": "What is the maximum payment amount per transaction?",
            "expected_answer": "The maximum standard payment amount is $10,000.00 per transaction.",
            "source": "payment_policy.pdf",
            "type": "table"
        },

        {
            "id": "Q05",
            "question": "What is the maximum refund amount?",
            "expected_answer": "The maximum refund amount is $5,000.00.",
            "source": "payment_policy.pdf",
            "type": "table"
        },

        {
            "id": "Q06",
            "question": "What authentication method does the Payment API use?",
            "expected_answer": "The Payment API uses Bearer Token authentication.",
            "source": "payment_policy.pdf",
            "type": "text"
        },

        {
            "id": "Q07",
            "question": "Which services receive requests from the API Gateway?",
            "expected_answer": "The API Gateway routes requests to the Payment Service and Order Service.",
            "source": "architecture.png",
            "type": "image"
        },

        {
            "id": "Q08",
            "question": "What is the price of the Wireless Keyboard?",
            "expected_answer": "The Wireless Keyboard costs $49.99 USD.",
            "source": "product_catalog.csv",
            "type": "structured"
        },

        {
            "id": "Q09",
            "question": "Which product has the highest price?",
            "expected_answer": "The 27-inch Monitor has the highest price at $249.99 USD.",
            "source": "product_catalog.csv",
            "type": "structured"
        },

        {
            "id": "Q10",
            "question": "How many USB-C Hubs are in stock?",
            "expected_answer": "There are 85 USB-C Hubs in stock.",
            "source": "product_catalog.csv",
            "type": "structured"
        },

        {
            "id": "Q11",
            "question": "What is Alice Johnson's membership level?",
            "expected_answer": "Alice Johnson has Gold membership.",
            "source": "customers.json",
            "type": "structured"
        },

        {
            "id": "Q12",
            "question": "How many orders has Carol Williams placed?",
            "expected_answer": "Carol Williams has placed 32 orders.",
            "source": "customers.json",
            "type": "structured"
        },

        {
            "id": "Q13",
            "question": "Who has Platinum membership?",
            "expected_answer": "Carol Williams has Platinum membership.",
            "source": "customers.json",
            "type": "structured"
        },

        {
            "id": "Q14",
            "question": "When was NovaTech Solutions founded?",
            "expected_answer": "NovaTech Solutions was founded in 2018.",
            "source": "company_info.html",
            "type": "html"
        },

        {
            "id": "Q15",
            "question": "Where is NovaTech Solutions headquartered?",
            "expected_answer": "NovaTech Solutions is headquartered in Bengaluru, India.",
            "source": "company_info.html",
            "type": "html"
        },

        {
            "id": "Q16",
            "question": "What are NovaTech's primary products?",
            "expected_answer": "NovaCloud, NovaAnalytics, and NovaSecure.",
            "source": "company_info.html",
            "type": "html"
        },

        {
            "id": "Q17",
            "question": "What was the company's revenue in 2024?",
            "expected_answer": "The company's revenue in 2024 was $130 million.",
            "source": "financial_report.xlsx",
            "type": "chart/table"
        },

        {
            "id": "Q18",
            "question": "Which department had the highest 2025 cost?",
            "expected_answer": "Engineering had the highest 2025 cost at $4.2 million.",
            "source": "financial_report.xlsx",
            "type": "table"
        },

        {
            "id": "Q19",
            "question": "What was the company's revenue in 2025?",
            "expected_answer": "The company's revenue in 2025 was $150 million.",
            "source": "financial_report.xlsx",
            "type": "chart/table"
        },

        {
            "id": "Q20",
            "question": "Which customer has Platinum membership and how many orders has that customer placed?",
            "expected_answer": "Carol Williams has Platinum membership and has placed 32 orders.",
            "source": "customers.json",
            "type": "reasoning"
        },

        {
            "id": "Q21",
            "question": "What is the maximum payment amount and where is NovaTech headquartered?",
            "expected_answer": "The maximum standard payment amount is $10,000.00 per transaction, and NovaTech is headquartered in Bengaluru, India.",
            "source": "payment_policy.pdf + company_info.html",
            "type": "cross-document"
        },

        {
            "id": "Q22",
            "question": "What is the most expensive product and what was revenue in 2025?",
            "expected_answer": "The 27-inch Monitor is the most expensive product at $249.99, and 2025 revenue was $150 million.",
            "source": "product_catalog.csv + financial_report.xlsx",
            "type": "cross-document"
        },

        {
            "id": "Q23",
            "question": "What is NovaTech Solutions' annual revenue in 2026?",
            "expected_answer": "The information is not provided in the documents.",
            "source": "none",
            "type": "unanswerable"
        },

        {
            "id": "Q24",
            "question": "Who is the CEO of NovaTech Solutions?",
            "expected_answer": "The information is not provided in the documents.",
            "source": "none",
            "type": "unanswerable"
        }

    ]

    data = {
        "description": "Multimodal RAG evaluation dataset",
        "questions": questions,
    }

    path.write_text(
        json.dumps(
            data,
            indent=2,
        ),
        encoding="utf-8",
    )

    print(f"Created: {path}")


# ============================================================
# 10. README
# ============================================================

def create_readme():

    path = OUTPUT_DIR / "README.md"

    content = """# Multimodal RAG Test Dataset

This dataset is designed to test a RAG pipeline against multiple
document types and content types.

## Files

- employee_handbook.docx
  - Text
  - Table
  - Embedded image

- payment_policy.pdf
  - Text
  - Table
  - Architecture diagram

- product_catalog.csv
  - Structured tabular data

- customers.json
  - Structured JSON data

- company_info.html
  - HTML text
  - Image reference

- financial_report.xlsx
  - Spreadsheet table
  - Bar chart

- architecture.png
  - Architecture diagram

- rag_test_questions.json
  - Direct questions
  - Table questions
  - Image questions
  - Chart questions
  - Cross-document questions
  - Unanswerable questions

## Suggested metadata

Your ingestion pipeline should preserve metadata such as:

source
file_type
page
sheet
section
element_type
table_id
image_id
chunk_id

## Important

After copying these documents into your RAG document directory,
run your ingestion process before starting the chatbot.

Example:

python ingest.py

Then:

python main.py
"""

    path.write_text(
        content,
        encoding="utf-8",
    )

    print(f"Created: {path}")


# ============================================================
# 11. ZIP
# ============================================================

def create_zip():

    zip_path = BASE_DIR / "rag_test_dataset.zip"

    with zipfile.ZipFile(
        zip_path,
        "w",
        zipfile.ZIP_DEFLATED,
    ) as zip_file:

        for file in OUTPUT_DIR.iterdir():

            if file.is_file():

                zip_file.write(
                    file,
                    arcname=file.name,
                )

    print(f"\nCreated ZIP: {zip_path}")


# ============================================================
# MAIN
# ============================================================

def main():

    print()
    print("=" * 60)
    print("Generating Multimodal RAG Test Dataset")
    print("=" * 60)
    print()

    create_architecture_png()

    create_remote_work_image()

    create_employee_handbook()

    create_payment_policy()

    create_product_catalog()

    create_customers()

    create_company_html()

    create_financial_report()

    create_questions()

    create_readme()

    create_zip()

    print()
    print("=" * 60)
    print("DONE")
    print("=" * 60)
    print()
    print(f"Dataset folder:")
    print(OUTPUT_DIR)
    print()
    print("ZIP:")
    print(BASE_DIR / "rag_test_dataset.zip")
    print()


if __name__ == "__main__":
    main()