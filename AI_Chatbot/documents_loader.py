import json
from pypdf import PdfReader
import pdfplumber

from bs4 import BeautifulSoup
from docx import Document
from pptx import Presentation
from openpyxl import load_workbook

def load_documents(self):
    """Load supported documents into a common representation."""

    if not self.input_file.is_dir():
        with open(
            self.input_file,
            "r",
            encoding="utf-8"
        ) as f:
            return json.load(f)

    documents = []

    for file_path in sorted(self.input_file.iterdir()):

        if not file_path.is_file():
            continue

        suffix = file_path.suffix.lower()

        try:
            if suffix == ".pdf":
                documents.extend(load_pdf(self, file_path))

            elif suffix == ".txt":
                documents.extend(load_txt(self, file_path))

            elif suffix == ".docx":
                documents.extend(load_docx(self, file_path))

            elif suffix == ".md":
                documents.extend(load_txt(self, file_path))

            elif suffix == ".csv":
                documents.extend(load_csv(self, file_path))

            elif suffix == ".json":
                documents.extend(load_json(self, file_path))

            elif suffix in [".html", ".htm"]:
                documents.extend(load_html(self, file_path))

            elif suffix == ".xlsx":
                documents.extend(load_xlsx(self, file_path))

            elif suffix == ".pptx":
                documents.extend(load_pptx(self, file_path))

            else:
                print(
                    f"Skipping unsupported file: "
                    f"{file_path.name}"
                )

        except Exception as e:
            print(
                f"Error loading {file_path.name}: {e}"
            )

    return documents

def load_pdf(self, file_path):

    documents = []

    reader = PdfReader(str(file_path))
    pending_table = None

    with pdfplumber.open(str(file_path)) as pdf:
        for page_number, page in enumerate(
            reader.pages,
            start=1
        ):
            text = page.extract_text() or ""

            if text.strip():
                documents.append({
                    "text": text,
                    "metadata": {
                        "source": file_path.name,
                        "file_type": "pdf",
                        "page": page_number
                    }
                })

            tables = pdf.pages[page_number - 1].extract_tables()

            for table_index, table in enumerate(tables):
                rows = [
                    [clean_table_cell(cell) for cell in row]
                    for row in table
                    if row and any(clean_table_cell(cell) for cell in row)
                ]

                if not rows:
                    continue

                column_count = max(len(row) for row in rows)

                if (
                    pending_table
                    and pending_table["page_end"] == page_number - 1
                    and pending_table["signature"] == column_count
                ):
                    pending_table["rows"].extend(rows)
                    pending_table["page_end"] = page_number
                    continue

                if pending_table:
                    documents.append(format_table_document(file_path, pending_table))

                pending_table = {
                    "rows": rows,
                    "page_start": page_number,
                    "page_end": page_number,
                    "table_index": table_index,
                    "signature": column_count,
                }

    if pending_table:
        documents.append(format_table_document(file_path, pending_table))

    return documents


def clean_table_cell(cell):
    return " ".join(str(cell or "").split())


def format_table_document(file_path, table):
    rows = []
    headers = table["rows"][0]

    for row in table["rows"]:
        cells = row + [""] * (len(headers) - len(row))
        rows.append(
            " | ".join(
                f"{header or f'Column {index + 1}'}: {cells[index]}"
                for index, header in enumerate(headers)
                if cells[index]
            )
        )

    return {
        "text": "Table\n\n" + "\n".join(rows),
        "metadata": {
            "source": file_path.name,
            "file_type": "pdf",
            "page": table["page_start"],
            "page_start": table["page_start"],
            "page_end": table["page_end"],
            "table_index": table["table_index"],
        }
    }

def load_txt(self, file_path):

    text = file_path.read_text(
        encoding="utf-8"
    )

    if not text.strip():
        return []

    return [{
        "text": text,
        "metadata": {
            "source": file_path.name,
            "file_type": file_path.suffix.lower().replace(".", ""),
            "page": None
        }
    }]

def load_docx(self, file_path):

    doc = Document(str(file_path))

    parts = []

    for paragraph in doc.paragraphs:

        text = paragraph.text.strip()

        if text:
            parts.append(text)

    # Include tables
    for table in doc.tables:

        for row in table.rows:

            cells = [
                cell.text.strip()
                for cell in row.cells
            ]

            row_text = " | ".join(
                cell for cell in cells
                if cell
            )

            if row_text:
                parts.append(row_text)

    text = "\n".join(parts)

    if not text.strip():
        return []

    return [{
        "text": text,
        "metadata": {
            "source": file_path.name,
            "file_type": "docx",
            "page": None
        }
    }]

def load_csv(self, file_path):

    import csv

    documents = []

    with open(
        file_path,
        "r",
        encoding="utf-8",
        newline=""
    ) as f:

        reader = csv.DictReader(f)

        for row_number, row in enumerate(
            reader,
            start=1
        ):

            parts = []

            for key, value in row.items():

                if value is None:
                    continue

                value = str(value).strip()

                if value:
                    parts.append(
                        f"{key}: {value}"
                    )

            text = "\n".join(parts)

            if not text.strip():
                continue

            documents.append({
                "text": text,
                "metadata": {
                    "source": file_path.name,
                    "file_type": "csv",
                    "page": None,
                    "row": row_number
                }
            })

    return documents

def load_json(self, file_path):

    with open(
        file_path,
        "r",
        encoding="utf-8"
    ) as f:

        data = json.load(f)

    text = json.dumps(
        data,
        indent=2,
        ensure_ascii=False
    )

    if not text.strip():
        return []

    return [{
        "text": text,
        "metadata": {
            "source": file_path.name,
            "file_type": "json",
            "page": None
        }
    }]

def load_html(self, file_path):

    html = file_path.read_text(
        encoding="utf-8"
    )

    soup = BeautifulSoup(
        html,
        "html.parser"
    )

    # Remove things that shouldn't enter RAG
    for element in soup([
        "script",
        "style",
        "nav",
        "footer",
        "header"
    ]):
        element.decompose()

    text = soup.get_text(
        separator="\n"
    )

    text = "\n".join(
        line.strip()
        for line in text.splitlines()
        if line.strip()
    )

    if not text.strip():
        return []

    return [{
        "text": text,
        "metadata": {
            "source": file_path.name,
            "file_type": "html",
            "page": None
        }
    }]

def load_xlsx(self, file_path):

    workbook = load_workbook(
        filename=file_path,
        read_only=True,
        data_only=True
    )

    documents = []

    for worksheet in workbook.worksheets:

        rows = []

        for row in worksheet.iter_rows(
            values_only=True
        ):

            values = [
                str(value).strip()
                for value in row
                if value is not None
            ]

            if values:
                rows.append(
                    " | ".join(values)
                )

        if not rows:
            continue

        text = "\n".join(rows)

        documents.append({
            "text": text,
            "metadata": {
                "source": file_path.name,
                "file_type": "xlsx",
                "page": None,
                "sheet": worksheet.title
            }
        })

    return documents

def load_pptx(self, file_path):

    presentation = Presentation(
        str(file_path)
    )

    documents = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1
    ):

        texts = []

        for shape in slide.shapes:

            if not hasattr(shape, "text"):
                continue

            text = shape.text.strip()

            if text:
                texts.append(text)

        text = "\n".join(texts)

        if not text.strip():
            continue

        documents.append({
            "text": text,
            "metadata": {
                "source": file_path.name,
                "file_type": "pptx",
                "page": slide_number,
                "slide": slide_number
            }
        })

    return documents




